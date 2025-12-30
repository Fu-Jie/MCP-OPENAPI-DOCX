"""Unit tests for table handler."""

import pytest
from docx import Document

from src.handlers.table_handler import TableHandler


class TestTableHandler:
    """Test cases for TableHandler class."""

    def setup_method(self):
        """Set up test fixtures."""
        self.handler = TableHandler()
        self.doc = Document()

    def test_create_table(self):
        """Test creating a table."""
        table = self.handler.create_table(self.doc, rows=3, cols=3)
        assert table is not None
        assert len(table.rows) == 3
        assert len(table.columns) == 3

    def test_create_table_with_data(self):
        """Test creating a table with data."""
        data = [
            ["A1", "B1", "C1"],
            ["A2", "B2", "C2"],
        ]
        table = self.handler.create_table(self.doc, rows=2, cols=3, data=data)

        assert table.cell(0, 0).text == "A1"
        assert table.cell(0, 2).text == "C1"
        assert table.cell(1, 1).text == "B2"

    def test_get_table_data(self):
        """Test getting table data."""
        data = [["R1C1", "R1C2"], ["R2C1", "R2C2"]]
        self.handler.create_table(self.doc, rows=2, cols=2, data=data)

        table_data = self.handler.get_table_data(self.doc, 0)
        assert table_data["rows"] == 2
        assert table_data["cols"] == 2
        assert table_data["data"][0][0] == "R1C1"

    def test_update_cell(self):
        """Test updating a cell value."""
        self.handler.create_table(self.doc, rows=2, cols=2)
        self.handler.update_cell(self.doc, 0, row=0, col=0, value="Updated")

        table_data = self.handler.get_table_data(self.doc, 0)
        assert table_data["data"][0][0] == "Updated"

    def test_add_row(self):
        """Test adding a row to a table."""
        self.handler.create_table(self.doc, rows=2, cols=2)
        self.handler.add_row(self.doc, 0, values=["New1", "New2"])

        table_data = self.handler.get_table_data(self.doc, 0)
        assert table_data["rows"] == 3

    def test_add_column(self):
        """Test adding a column to a table."""
        self.handler.create_table(self.doc, rows=2, cols=2)
        self.handler.add_column(self.doc, 0)

        table_data = self.handler.get_table_data(self.doc, 0)
        assert table_data["cols"] == 3

    def test_delete_row(self):
        """Test deleting a row from a table."""
        self.handler.create_table(self.doc, rows=3, cols=2)
        self.handler.delete_row(self.doc, 0, row_index=1)

        table_data = self.handler.get_table_data(self.doc, 0)
        assert table_data["rows"] == 2

    def test_list_tables(self):
        """Test listing all tables."""
        self.handler.create_table(self.doc, rows=2, cols=2)
        self.handler.create_table(self.doc, rows=3, cols=3)

        tables = self.handler.list_tables(self.doc)
        assert len(tables) == 2
        assert tables[0]["rows"] == 2
        assert tables[1]["rows"] == 3

    def test_merge_cells(self):
        """Test merging cells."""
        self.handler.create_table(self.doc, rows=3, cols=3)
        self.handler.merge_cells(
            self.doc, 0,
            start_row=0, start_col=0,
            end_row=0, end_col=1
        )

        # Verify merge happened (no exception means success)
        assert len(self.doc.tables) == 1
